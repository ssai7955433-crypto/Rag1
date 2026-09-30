from __future__ import annotations

import csv
from dataclasses import dataclass
from io import BytesIO, StringIO
import json
from pathlib import Path
import re
from html.parser import HTMLParser

import fitz
from docx import Document


@dataclass
class LoadedDocument:
    filename: str
    source_type: str
    pages: list[tuple[int | None, str]]

    @property
    def text(self) -> str:
        return "\n\n".join(text for _, text in self.pages)

    @property
    def page_count(self) -> int:
        return sum(1 for _, text in self.pages if text.strip())


class _TextExtractor(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []

    def handle_data(self, data: str) -> None:
        if data.strip():
            self.parts.append(data.strip())


def _plain_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, (list, tuple)):
        return "; ".join(filter(None, (_plain_text(item) for item in value)))
    if isinstance(value, dict):
        return "; ".join(
            f"{key}: {text}" for key, item in value.items() if (text := _plain_text(item))
        )
    text = str(value).strip()
    if "<" in text and ">" in text:
        parser = _TextExtractor()
        try:
            parser.feed(text)
            extracted = " ".join(parser.parts)
            if extracted:
                return re.sub(r"\s+", " ", extracted).strip()
        except Exception:
            pass
    return re.sub(r"\s+", " ", text).strip()


def _format_product(product: dict, index: int) -> str:
    values = {str(key).strip().lower().replace(" ", "_"): value for key, value in product.items()}
    aliases = {
        "Product": ("product_name", "title", "name", "product"),
        "SKU": ("sku", "product_id", "id"),
        "Brand": ("brand", "vendor", "maker"),
        "Category": ("category", "product_type", "type"),
        "Price": ("price", "price_amount", "sale_price"),
        "Currency": ("currency",),
        "Availability": ("availability", "inventory_status", "stock_status", "in_stock"),
        "Average rating": ("average_rating", "rating", "stars", "review_rating"),
        "Features": ("features", "specifications", "specs"),
        "Description": ("description", "product_description", "body_html", "details"),
        "Customer reviews": ("reviews", "customer_reviews", "review", "review_text", "review_body", "body"),
    }
    product_keys = aliases["Product"]
    product_name = next((values[key] for key in product_keys if values.get(key)), f"Catalog item {index}")
    lines = [f"Product: {_plain_text(product_name)}"]
    consumed = {key for keys in aliases.values() for key in keys}
    for label, keys in aliases.items():
        if label == "Product":
            continue
        value = next((values[key] for key in keys if values.get(key) not in (None, "")), None)
        if value is not None:
            rendered = _plain_text(value)
            if rendered:
                lines.append(f"{label}: {rendered}")
    for key, value in values.items():
        if key not in consumed and value not in (None, ""):
            rendered = _plain_text(value)
            if rendered:
                lines.append(f"{key.replace('_', ' ').title()}: {rendered}")
    return "\n".join(lines)


def _load_catalog(file_bytes: bytes, filename: str, suffix: str) -> LoadedDocument:
    if suffix == ".csv":
        text_stream = StringIO(file_bytes.decode("utf-8-sig"), newline="")
        rows = list(csv.DictReader(text_stream))
        source_type = "csv_catalog"
    else:
        payload = json.loads(file_bytes.decode("utf-8-sig"))
        if isinstance(payload, dict) and isinstance(payload.get("products"), list):
            rows = payload["products"]
        elif isinstance(payload, list):
            rows = payload
        elif isinstance(payload, dict):
            rows = [payload]
        else:
            rows = []
        source_type = "json_catalog"

    pages: list[tuple[int | None, str]] = []
    for item_number, row in enumerate(rows, start=1):
        if isinstance(row, dict) and any(_plain_text(value) for value in row.values()):
            text = _format_product(row, item_number)
            if text.strip():
                pages.append((item_number, text))
    return LoadedDocument(filename, source_type, pages)


def load_document(file_bytes: bytes, filename: str) -> LoadedDocument:
    suffix = Path(filename).suffix.lower()
    if suffix in {".csv", ".json"}:
        return _load_catalog(file_bytes, filename, suffix)
    if suffix == ".pdf":
        pages: list[tuple[int | None, str]] = []
        with fitz.open(stream=file_bytes, filetype="pdf") as document:
            for page_number, page in enumerate(document, start=1):
                text = page.get_text("text")
                if text and text.strip():
                    pages.append((page_number, text))
        return LoadedDocument(filename, "pdf", pages)

    if suffix == ".docx":
        document = Document(BytesIO(file_bytes))
        parts = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            for row in table.rows:
                values = [cell.text.strip() for cell in row.cells]
                row_text = " | ".join(value for value in values if value)
                if row_text:
                    parts.append(row_text)
        return LoadedDocument(filename, "docx", [(None, "\n\n".join(parts))] if parts else [])

    raise ValueError("Unsupported file type. Upload a CSV, JSON, PDF, or DOCX file.")
